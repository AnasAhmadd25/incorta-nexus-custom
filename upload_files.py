#  i want to make a here a logic to read file uploads, detect the type, and translate it as text and add it to the chat

import csv
import filetype
from PyPDF2 import PdfReader
from abc import ABC, abstractmethod
from docx import Document
import pandas as pd
import openpyxl
from typing import List
import os
# import easyocr  # Commented out to avoid heavy dependencies
import logging
# ----------------- Main class -----------------
class FileHandler(ABC):
    def __init__(self, file_path):
        self.file_path = file_path
        self.file_type = self.detect_file_type()

    def detect_file_type(self):
        kind = filetype.guess(self.file_path)
        if kind is None:
            return "Unknown"
        return kind.mime

    @abstractmethod
    def get_text(self):
        pass

    @abstractmethod
    def add_text(self, text: str):
        pass

# ---------- each type class ----------


# Csv to txt 

class CsvHandler(FileHandler):
    def get_text(self): # need to make it robust to more encodings
        with open(self.file_path, 'r', encoding='utf-8') as file:
            return file.read()
    
    def add_text(self, text: str):
        return text
            



class PdfHandler(FileHandler):
    def get_text(self):
        with open(self.file_path, 'rb') as file:
            reader = PdfReader(file)
            text = ''
            for page_num, page in enumerate(reader.pages, 1):
                text += page.extract_text()
                # attach page number with a divider
                text += f"\n--- Page {page_num} ---\n"
            return text
    
    def add_text(self, text: str):
        return text


class DocxHandler(FileHandler):
    def get_text(self):
        doc = Document(self.file_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    
    def add_text(self, text: str):
        return text

class TxtHandler(FileHandler):
    def get_text(self) -> str:
        with open(self.file_path, 'r', encoding='utf-8') as file:
            return file.read()
    
    def add_text(self, text: str):
        return text

       

class ExcelHandler(FileHandler):
    def get_text(self, max_rows_per_sheet: int = 1000) -> str:
        try:
    
            sheet_names = self.get_sheet_names()
            
            text_parts = []
            
            for sheet_name in sheet_names:
                try:

                    df = pd.read_excel(
                        self.file_path, 
                        sheet_name=sheet_name,
                        nrows=max_rows_per_sheet
                    )
                    
                    if df.empty:
                        sheet_text = f"This is a sheet named '{sheet_name}' data: [Empty sheet]"
                    else:
                        metadata = f"Sheet '{sheet_name}' contains {len(df)} rows and {len(df.columns)} columns."
                        if len(df) == max_rows_per_sheet:
                            metadata += f" (Limited to first {max_rows_per_sheet} rows)"
                        
                        # Convert to text
                        data_text = df.to_string(index=False)
                        sheet_text = f"This is a sheet named '{sheet_name}' data:\n{metadata}\n\n{data_text}"
                    
                    text_parts.append(sheet_text)
                    
                except Exception as e:
                    text_parts.append(f"This is a sheet named '{sheet_name}' data: [Error reading sheet: {str(e)}]")
            
            return "\n\n" + "="*50 + "\n\n".join(text_parts)
            
        except Exception as e:
            return f"Error reading Excel file: {str(e)}"
    
    def get_sheet_names(self) -> List[str]:
        """Fast method to get sheet names using openpyxl"""
        try:
            workbook = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            return workbook.sheetnames
        except Exception as e:
            print(f"Error getting sheet names: {e}")
            return []
    
    def add_text(self, text: str) -> str:
        """Enhanced context addition"""
        context = """The following data comes from an Excel spreadsheet with multiple sheets.
Each sheet is clearly labeled with its name. Use this data to answer questions about the spreadsheet content.

"""
        return context + text

# class ImageHandler(FileHandler):
#     def __init__(self, file_path):
#         super().__init__(file_path)
#         import easyocr
#         logging.getLogger('easyocr').setLevel(logging.ERROR)
#         self.reader = easyocr.Reader(['en'])
#
#     def get_text(self):
#         try:
#             result = self.reader.readtext(self.file_path, detail=0)
#             return '\n'.join(result)
#         except Exception as e:
#             return f"Error reading image file: {str(e)}"
#
#     def add_text(self, text: str):
#         return text
# factory pattern 


class FileHandlerFactory:
    @staticmethod
    def create_handler(file_path: str) -> FileHandler:
        kind = filetype.guess(file_path)
        if kind is not None:
            mime_type = kind.mime
        else:
            mime_type = FileHandlerFactory._detect_by_extension(file_path)

        switcher = {
            "text/csv": CsvHandler(file_path),
            "application/pdf": PdfHandler(file_path),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocxHandler(file_path),
            "text/plain": TxtHandler(file_path),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ExcelHandler(file_path),
            "application/vnd.ms-excel": ExcelHandler(file_path),
            # Image handlers commented out to avoid easyocr dependency
            # "image/png": ImageHandler(file_path),
            # "image/jpeg": ImageHandler(file_path),
            # "image/jpg": ImageHandler(file_path)
        }

        handler = switcher.get(mime_type)
        if handler is None:
            _, ext = os.path.splitext(file_path)
            raise ValueError(f"Unsupported file type: {mime_type} (file extension: {ext}). Supported types: PDF, DOCX, TXT, CSV, XLSX, XLS")
        return handler
    
    @staticmethod
    def _detect_by_extension(file_path: str) -> str:
        """Fallback method to detect file type by extension"""
        _, ext = os.path.splitext(file_path.lower())
        
        extension_map = {
            '.txt': 'text/plain',
            '.csv': 'text/csv',
            '.pdf': 'application/pdf',
            '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.xls': 'application/vnd.ms-excel'
        }
        
        return extension_map.get(ext, "unknown")
